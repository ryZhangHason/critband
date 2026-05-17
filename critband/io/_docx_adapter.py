"""DOCX adapter using python-docx.

Extracts tables from Word documents.
"""

from typing import Dict, List

import numpy as np

from critband.io._common import filter_numerical_columns, try_parse_float


def read_docx(path) -> Dict[str, Dict[str, np.ndarray]]:
    """Read DOCX file into {sheet_name: {column_name: array}}."""
    import docx

    doc = docx.Document(path)
    return _parse_document(doc)


def read_docx_buffer(buf) -> Dict[str, Dict[str, np.ndarray]]:
    """Read DOCX from a buffer."""
    import docx

    doc = docx.Document(buf)
    return _parse_document(doc)


def _parse_document(doc) -> Dict[str, Dict[str, np.ndarray]]:
    tables: Dict[str, Dict[str, np.ndarray]] = {}
    for i, table in enumerate(doc.tables):
        rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        if len(rows) < 2:
            continue

        headers = rows[0]
        if not headers:
            continue

        columns: Dict[str, List] = {h: [] for h in headers}
        for row in rows[1:]:
            for j, h in enumerate(headers):
                cell_val = row[j] if j < len(row) else ""
                columns[h].append(try_parse_float(cell_val))

        result = filter_numerical_columns(columns)
        if result:
            label = f"Table {i + 1}"
            tables[label] = result

    return tables

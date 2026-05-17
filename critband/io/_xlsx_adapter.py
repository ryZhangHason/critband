"""XLSX adapter using openpyxl."""

from typing import Dict, List

import numpy as np

from critband.io._common import filter_numerical_columns, try_parse_float


def read_xlsx(path) -> Dict[str, Dict[str, np.ndarray]]:
    """Read XLSX file into {sheet_name: {column_name: array}}."""
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    return _parse_workbook(wb)


def read_xlsx_buffer(buf) -> Dict[str, Dict[str, np.ndarray]]:
    """Read XLSX from a buffer."""
    import openpyxl

    wb = openpyxl.load_workbook(buf, read_only=True, data_only=True)
    return _parse_workbook(wb)


def _parse_workbook(wb) -> Dict[str, Dict[str, np.ndarray]]:
    tables: Dict[str, Dict[str, np.ndarray]] = {}
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue

        # First row is headers
        headers = [str(c) if c is not None else f"Col{i + 1}" for i, c in enumerate(rows[0])]
        if not headers:
            continue

        columns: Dict[str, List] = {h: [] for h in headers}
        for row in rows[1:]:
            if row is None:
                continue
            for i, h in enumerate(headers):
                val = row[i] if i < len(row) else None
                if val is not None and not isinstance(val, str):
                    try:
                        val = float(val)
                    except (ValueError, TypeError):
                        val = None
                elif isinstance(val, str):
                    val = try_parse_float(val)
                columns[h].append(val)

        result = filter_numerical_columns(columns)
        if result:
            tables[ws.title] = result

    wb.close()
    return tables
